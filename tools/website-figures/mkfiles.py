# Sample files for the website's speed figures (deterministic).
import random, zipfile, sys
from openpyxl import Workbook
out = sys.argv[1]
random.seed(7)
regions = ["Norte", "Sur", "Este", "Oeste", "Centro"]
products = ["Portátil", "Monitor", "Teclado", "Ratón", "Impresora", "Tableta", "Auriculares", "Cámara"]
wb = Workbook(write_only=True)
ws = wb.create_sheet("Ventas")
ws.append(["Pedido", "Fecha", "Región", "Producto", "Unidades", "Precio", "Descuento", "Importe", "IVA", "Total"])
for i in range(2, 100002):
    ws.append([100000 + i, f"2026-{random.randint(1,12):02d}-{random.randint(1,28):02d}", random.choice(regions),
               random.choice(products), random.randint(1, 40), round(random.uniform(5, 1500), 2),
               round(random.choice([0, 0, .05, .1, .15]), 2), f"=E{i}*F{i}*(1-G{i})", f"=H{i}*0.21", f"=H{i}+I{i}"])
wb.save(f"{out}/big.xlsx")
words = ("el informe trimestral recoge las ventas por región y producto con un análisis de tendencias "
         "costes márgenes previsiones y riesgos que la dirección revisa cada mes para decidir inversiones").split()
paras = []
for s in range(1, 121):
    paras.append(f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Sección {s}</w:t></w:r></w:p>')
    for p in range(12):
        text = " ".join(random.choice(words) for _ in range(90)).capitalize() + "."
        bold = '<w:rPr><w:b/></w:rPr>' if p == 0 else ''
        paras.append(f'<w:p><w:r>{bold}<w:t xml:space="preserve">{text}</w:t></w:r></w:p>')
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
doc = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body>{"".join(paras)}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr></w:body></w:document>'
styles = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles {W}><w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:rPr><w:sz w:val="22"/></w:rPr></w:style><w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:pPr><w:keepNext/><w:spacing w:before="240" w:after="120"/><w:outlineLvl w:val="0"/></w:pPr><w:rPr><w:b/><w:sz w:val="32"/></w:rPr></w:style></w:styles>'
ct = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/><Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>'
rels = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'
drels = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>'
with zipfile.ZipFile(f"{out}/big.docx", "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("[Content_Types].xml", ct); z.writestr("_rels/.rels", rels)
    z.writestr("word/document.xml", doc); z.writestr("word/styles.xml", styles); z.writestr("word/_rels/document.xml.rels", drels)

# Everyday files: a two-page letter and a small monthly budget.
letter = ["Estimada Sra. García:"] + [" ".join(random.choice(words) for _ in range(70)).capitalize() + "." for _ in range(9)] + ["Atentamente,", "Laura Martín"]
body = "".join(f'<w:p><w:r><w:t xml:space="preserve">{t}</w:t></w:r></w:p>' for t in letter)
doc = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Propuesta de colaboración</w:t></w:r></w:p>{body}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr></w:body></w:document>'
with zipfile.ZipFile(f"{out}/carta.docx", "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("[Content_Types].xml", ct); z.writestr("_rels/.rels", rels)
    z.writestr("word/document.xml", doc); z.writestr("word/styles.xml", styles); z.writestr("word/_rels/document.xml.rels", drels)
wb = Workbook()
ws = wb.active; ws.title = "Presupuesto"
ws.append(["Concepto", "Categoría", "Ene", "Feb", "Mar", "Abr", "May", "Jun", "Total"])
cats = ["Oficina", "Personal", "Suministros", "Viajes", "Marketing"]
for i in range(2, 42):
    ws.append([f"Partida {i - 1}", random.choice(cats)] + [round(random.uniform(50, 3000), 2) for _ in range(6)] + [f"=SUM(C{i}:H{i})"])
ws.append(["Total", ""] + [f"=SUM({c}2:{c}41)" for c in "CDEFGHI"])
wb.save(f"{out}/presupuesto.xlsx")
